import pandas as pd
import openpyxl

df = pd.read_parquet("Libro_Gestion.parquet")
print("=== PARQUET ===")
print(df.dtypes)
print(df.head(2).to_string())

wb = openpyxl.load_workbook("Robot 2026.xlsx")
print("\n=== EXCEL SHEETS ===")
for sh in wb.sheetnames:
    ws = wb[sh]
    print(f"\nHoja: {sh}")
    for row in ws.iter_rows(min_row=1, max_row=3, values_only=True):
        print(row)
