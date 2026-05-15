"""
Endpoints for the KPI mapping configuration UI.
"""
import json
from fastapi import APIRouter, HTTPException
import openpyxl

from ..config import settings
from ..services import parquet_service as pq
from ..services import mapping_store
from ..services.mapper_service import get_excel_kpis_for_setup

router = APIRouter(prefix="/api/setup", tags=["setup"])

KPI_STRUCTURE_FILE = settings.parquet_path.parent / "kpi_structure.json"


def _get_df():
    if not settings.parquet_path.exists():
        raise HTTPException(
            status_code=404,
            detail=f"Parquet no encontrado: {settings.parquet_path}",
        )
    return pq.get_cached_df(settings.parquet_path)


def _load_structure(company: str) -> list | None:
    """Load predefined KPI structure from JSON, enriching KPI rows with their section."""
    if not KPI_STRUCTURE_FILE.exists():
        return None
    data = json.loads(KPI_STRUCTURE_FILE.read_text(encoding="utf-8"))
    raw = data.get(company)
    if raw is None:
        # Company not in kpi_structure.json — try to build Costos Ajustados section only
        result: list = []
        _append_costos_aj(company, result)
        return result if result else None
    current_section = ""
    result = []
    for item in raw:
        if item["type"] == "header":
            current_section = item["label"]
            result.append(item)
        else:
            kpi_id = f"{current_section}||{item['label']}"
            result.append({**item, "section": current_section, "id": kpi_id})
    _append_costos_aj(company, result)
    return result


def _append_costos_aj(company: str, result: list) -> None:
    """Append Costos Ajustados section to structure, sourced from params_financieros.json."""
    params_path = KPI_STRUCTURE_FILE.parent / "params_financieros.json"
    if not params_path.exists():
        return
    fin = json.loads(params_path.read_text(encoding="utf-8"))
    grupos  = fin.get("costos_aj_grupos", {}).get(company, [])
    summary = fin.get("costos_aj_summary", [])
    if not grupos and not summary:
        return

    existing = mapping_store.load_mappings().get(company, {})

    def _has_auto_match(label: str) -> bool:
        """True if exactly one mapping outside Costos Ajustados resolves this label."""
        matches = [k for k in existing
                   if k.endswith(f"||{label}") and not k.startswith("Costos Ajustados")]
        return len(matches) == 1

    section = "Costos Ajustados"
    header_added = False
    seen_labels: set[str] = set()

    def _ensure_header():
        nonlocal header_added
        if not header_added:
            result.append({"type": "header", "label": section})
            header_added = True

    # Grupos: always show all items (user wants full visibility)
    for grp in grupos:
        grp_items = []
        sas = grp.get("subareas", [])
        # Fila ∑ Total solo si ya existe un mapeo directo a nivel de grupo (clave 2 o 3 partes)
        grp_total_id   = f"{section}||{grp['label']}"
        grp_total_id3  = f"{section}||{grp['label']}||Total"
        if len(sas) > 1 and (grp_total_id in existing or grp_total_id3 in existing):
            grp_items.append({
                "type": "kpi",
                "id": grp_total_id,
                "section": section,
                "code": grp["key"], "label": "Total", "unit": "kUS$", "row": None,
                "is_group_total": True,
            })
        for sa in sas:
            if sa["label"] in seen_labels:
                continue
            seen_labels.add(sa["label"])
            grp_items.append({
                "type": "kpi",
                "id": f"{section}||{grp['label']}||{sa['label']}",
                "section": section,
                "code": sa["key"], "label": sa["label"], "unit": "kUS$", "row": None,
            })
        if grp_items:
            _ensure_header()
            result.append({"type": "subheader", "label": grp["label"]})
            result.extend(grp_items)

    # Summary items: skip those already uniquely resolved from another section
    sum_items = []
    for sr in summary:
        if sr["label"] in seen_labels:
            continue
        seen_labels.add(sr["label"])
        if _has_auto_match(sr["label"]):
            continue
        sum_items.append({
            "type": "kpi",
            "id": f"{section}||{sr['label']}",
            "section": section,
            "code": sr["key"], "label": sr["label"], "unit": "kUS$", "row": None,
        })

    if sum_items:
        _ensure_header()
        result.append({"type": "subheader", "label": "Ajustes Financieros"})
        result.extend(sum_items)


@router.get("/excel-kpis/{company}")
def get_excel_kpis(company: str):
    """KPI structure for this company: uses kpi_structure.json if available,
    otherwise falls back to reading the Excel template."""
    structure = _load_structure(company)
    if structure is not None:
        return {"company": company, "kpis": structure}

    if not settings.template_path.exists():
        raise HTTPException(404, f"Template no encontrado: {settings.template_path}")
    wb = openpyxl.load_workbook(settings.template_path, read_only=True, data_only=True)
    kpis = get_excel_kpis_for_setup(wb, company)
    return {"company": company, "kpis": kpis}


@router.get("/parquet-kpis/{company}")
def get_parquet_kpis(company: str):
    """All unique (kpi_name, hoja) combinations in the parquet for this company.
    Returns list of {name, hoja, title} sorted by (hoja, name).
    A KPI that appears in multiple hojas generates one entry per hoja.
    name = 'subcategory||kpi' when subcategory exists, plain 'kpi' otherwise.
    """
    df = _get_df()
    company_df = df[df["compania"] == company]
    seen: set[tuple] = set()
    result: list[dict] = []
    if "kpi" in company_df.columns:
        has_sub   = "subcategory" in company_df.columns
        has_hoja  = "hoja"  in company_df.columns
        has_title = "title" in company_df.columns
        cols = ["kpi"]
        if has_sub:   cols.append("subcategory")
        if has_hoja:  cols.append("hoja")
        if has_title: cols.append("title")
        for _, row in company_df[cols].drop_duplicates().iterrows():
            kpi_name    = str(row.get("kpi", "")         or "").strip()
            subcategory = str(row.get("subcategory", "") or "").strip()
            hoja        = str(row.get("hoja",  "") or "").strip() if has_hoja  else ""
            title       = str(row.get("title", "") or "").strip() if has_title else ""
            if not kpi_name:
                continue
            name = f"{subcategory}||{kpi_name}" if subcategory else kpi_name
            key = (name, hoja)
            if key not in seen:
                seen.add(key)
                result.append({"name": name, "hoja": hoja, "title": title})
    result.sort(key=lambda x: (x["hoja"], x["name"]))
    return {"company": company, "kpis": result}


@router.get("/parquet-snapshot/{company}")
def get_parquet_snapshot(company: str, tipo: str, anio: int, mes: int):
    """Diagnostic: returns KPI names available for a specific tipo+periodo."""
    df = _get_df()
    periodo = anio * 100 + mes
    subset = df[(df["compania"] == company) & (df["tipo"] == tipo) & (df["periodo"] == periodo)]
    kpis: list[str] = []
    if "kpi" in subset.columns:
        for _, row in subset[["kpi", "subcategory"] if "subcategory" in subset.columns else ["kpi"]].drop_duplicates().iterrows():
            kpi_name = str(row.get("kpi", "") or "").strip()
            subcategory = str(row.get("subcategory", "") or "").strip()
            if not kpi_name:
                continue
            qualified = f"{subcategory}||{kpi_name}" if subcategory else kpi_name
            if qualified not in kpis:
                kpis.append(qualified)
    periodos_disponibles = sorted(
        df[(df["compania"] == company) & (df["tipo"] == tipo)]["periodo"].dropna().unique().tolist()
    )
    return {
        "company": company,
        "tipo": tipo,
        "periodo_buscado": periodo,
        "kpis_en_periodo": sorted(kpis),
        "periodos_disponibles": periodos_disponibles,
    }


def _migrate_mappings(mappings: dict) -> dict:
    """Upgrade old bare-label keys (no ||) to section||label format using kpi_structure.json."""
    if not KPI_STRUCTURE_FILE.exists():
        return mappings
    structure = json.loads(KPI_STRUCTURE_FILE.read_text(encoding="utf-8"))
    changed = False
    for company, items in structure.items():
        if company not in mappings:
            continue
        # Build label -> section lookup (first occurrence wins for unique labels)
        label_to_section: dict[str, str] = {}
        cur_sec = ""
        for item in items:
            if item["type"] == "header":
                cur_sec = item["label"]
            elif item["label"] not in label_to_section:
                label_to_section[item["label"]] = cur_sec
        old_maps = mappings[company]
        new_maps: dict = {}
        for key, val in old_maps.items():
            if "||" in key:
                new_maps[key] = val
            elif key in label_to_section:
                new_key = f"{label_to_section[key]}||{key}"
                new_maps[new_key] = val
                changed = True
            else:
                new_maps[key] = val
        mappings[company] = new_maps
    if changed:
        mapping_store.save_mappings(mappings)
    return mappings


@router.get("/mappings")
def get_mappings():
    mappings = mapping_store.load_mappings()
    return _migrate_mappings(mappings)


@router.post("/mappings")
def save_mappings(data: dict):
    mapping_store.save_mappings(data)
    return {"ok": True}


@router.get("/sheet-dump/{sheet_name}")
def dump_sheet(sheet_name: str):
    """Temporary: dumps a sheet's non-empty cell content as JSON."""
    if not settings.template_path.exists():
        raise HTTPException(404, "Template no encontrado")
    try:
        wb = openpyxl.load_workbook(settings.template_path, data_only=True)
    except Exception as e:
        raise HTTPException(500, f"Error abriendo Excel: {e}")
    if sheet_name not in wb.sheetnames:
        raise HTTPException(404, f"Hoja '{sheet_name}' no encontrada. Hojas: {wb.sheetnames}")
    ws = wb[sheet_name]
    rows = []
    try:
        for row in ws.iter_rows():
            cells = []
            for c in row:
                try:
                    if c.value is not None:
                        cells.append({"col": c.column, "val": str(c.value)})
                except Exception:
                    pass
            if cells:
                rows.append({"row": row[0].row, "cells": cells})
    except Exception as e:
        return {"sheet": sheet_name, "error": str(e), "rows": rows}
    return {"sheet": sheet_name, "rows": rows}
