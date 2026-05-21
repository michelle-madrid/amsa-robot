"""
Visualización de KPIs por rango de fechas (multi-año).
Soporta dos comparaciones con sus propios tipo/rango.
"""
import io
import json
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

from ..config import settings
from ..services import parquet_service as pq
from ..services.mapping_store import load_mappings
from ..services.mapper_service import get_excel_kpis_for_setup
from .setup import _append_costos_aj

router = APIRouter(prefix="/api/visualizar", tags=["visualizar"])

KPI_STRUCTURE_FILE   = settings.parquet_path.parent / "kpi_structure.json"
FORMULA_PARAMS_FILE  = settings.parquet_path.parent / "formula_params.json"
MONTH_FIELDS = pq.MONTH_FIELDS
MONTH_BY_NUM = {i+1: f for i, f in enumerate(MONTH_FIELDS)}


def _load_visualizar_scales(company: str) -> dict[str, float]:
    if FORMULA_PARAMS_FILE.exists():
        data = json.loads(FORMULA_PARAMS_FILE.read_text(encoding="utf-8"))
        return data.get("visualizar_scales_per_company", {}).get(company, {})
    return {}


_TIPO_ALIASES: list[set[str]] = [
    {"ppto", "plan", "presupuesto", "budget"},
    {"real", "actual"},
]


def _tipo_variants(tipo: str) -> set[str]:
    t = tipo.strip().lower()
    for group in _TIPO_ALIASES:
        if t in group:
            return group
    return {t}


def _apply_op(a, op, b):
    if op == "+":
        if a is None and b is None: return None
        return (a or 0.0) + (b or 0.0)
    if a is None or b is None:
        return None
    if op == "-": return a - b
    if op == "*": return a * b
    if op == "/" and b != 0: return a / b
    return None


def _apply_scale(val, mapped: dict):
    """Apply optional scalar: (A op B) scale_op scale."""
    if val is None:
        return None
    scale = mapped.get("scale")
    if scale is None:
        return val
    scale_op = mapped.get("scale_op", "*")
    return _apply_op(val, scale_op, float(scale))


def _empty_series(n: int) -> list:
    return [None] * n


def _build_columns(año_ini, mes_ini, año_fin, mes_fin) -> list[dict]:
    """List of {year, month, label} for every month in the range."""
    cols = []
    y, m = año_ini, mes_ini
    while (y, m) <= (año_fin, mes_fin):
        short = ["Ene","Feb","Mar","Abr","May","Jun","Jul","Ago","Sep","Oct","Nov","Dic"][m-1]
        cols.append({"year": y, "month": m, "label": f"{short}'{str(y)[2:]}"})
        m += 1
        if m > 12:
            m = 1
            y += 1
    return cols


def _fetch_year_snapshot(df_co, tipo: str, year: int, snap_mes: int) -> dict[str, dict]:
    """Returns {'subcategory||kpi': {field: value}} for the given snapshot period."""
    variants = _tipo_variants(tipo)
    df_tipo = df_co[df_co["tipo"].str.strip().str.lower().isin(variants)]

    periodo = year * 100 + snap_mes
    df = df_tipo[df_tipo["periodo"] == periodo]

    if df.empty:
        year_min = year * 100 + 1
        df_year = df_tipo[(df_tipo["periodo"] >= year_min) & (df_tipo["periodo"] <= periodo)]
        if not df_year.empty:
            periodo = int(df_year["periodo"].max())
            df = df_tipo[df_tipo["periodo"] == periodo]

    result: dict[str, dict] = {}
    cols = [c for c in pq.MONTH_FIELDS + ["mes", "ytd"] if c in df.columns]
    for row in df[["kpi", "subcategory"] + cols].itertuples(index=False):
        kpi_name    = str(row.kpi        or "").strip()
        subcategory = str(row.subcategory or "").strip()
        if not kpi_name:
            continue
        key = f"{subcategory}||{kpi_name}" if subcategory else kpi_name
        if key not in result:
            d: dict = {}
            for f in pq.MONTH_FIELDS:
                v = getattr(row, f, None)
                d[f] = float(v) if v is not None and str(v) != "nan" else None
            v = getattr(row, "mes", None)
            d["mes"] = float(v) if v is not None and str(v) != "nan" else None
            v = getattr(row, "ytd", None)
            d["ytd"] = float(v) if v is not None and str(v) != "nan" else None
            result[key] = d
    return result


def _resolve_snapshot(snapshot: dict[str, dict], src: str) -> dict:
    """Look up a source in the snapshot. Supports exact match or suffix match for plain names.
    If src is 'A||B' and exact match fails, also tries matching just 'B' (kpi without subcategory).
    """
    if src in snapshot:
        return snapshot[src]
    suffix = f"||{src}"
    matches = [v for k, v in snapshot.items() if k.endswith(suffix)]
    if len(matches) == 1:
        return matches[0]
    # If src contains "||" (e.g. "Item||Ingresos Molibdeno"), also try the kpi-only part
    if "||" in src:
        plain = src.split("||", 1)[1]
        if plain in snapshot:
            return snapshot[plain]
        suffix2 = f"||{plain}"
        matches2 = [v for k, v in snapshot.items() if k.endswith(suffix2)]
        if len(matches2) == 1:
            return matches2[0]
    return {}


def _get_series(snapshots: dict[int, dict], columns: list[dict], sources: list[str]) -> list:
    """Returns one value per column for the given sources (summed if multiple).
    snapshots must be pre-built: {year: snapshot_dict}.
    """
    if not sources:
        return _empty_series(len(columns))
    result = []
    for col in columns:
        yr, mo = col["year"], col["month"]
        field  = MONTH_BY_NUM[mo]
        total  = None
        snap   = snapshots.get(yr, {})
        for src in sources:
            val = _resolve_snapshot(snap, src).get(field)
            if val is not None:
                total = (total or 0.0) + val
        result.append(total)
    return result


def _derive_mes_ytd(vals: list, cols: list[dict], año_fin: int) -> tuple:
    """Derive mes and ytd from the already-computed column values.
    mes = last column value (None if that month has no data).
    ytd = sum of all column values within año_fin (skips None).
    """
    mes = vals[-1] if vals else None
    ytd_vals = [v for col, v in zip(cols, vals) if col["year"] == año_fin and v is not None]
    ytd = sum(ytd_vals) if ytd_vals else None
    return mes, ytd


def _get_company_items(company: str) -> list:
    """Carga la estructura KPI de la compañía desde kpi_structure.json.
    Si la compañía no tiene estructura definida, hace fallback al Excel template
    (igual que el tab Configuración), para que siempre se vea la estructura.
    """
    if KPI_STRUCTURE_FILE.exists():
        data = json.loads(KPI_STRUCTURE_FILE.read_text(encoding="utf-8"))
        items = data.get(company, [])
        if items:
            result = list(items)
            _append_costos_aj(company, result)
            return result
    # Fallback al template Excel
    if settings.template_path.exists():
        try:
            wb = openpyxl.load_workbook(settings.template_path, read_only=True, data_only=True)
            return get_excel_kpis_for_setup(wb, company)
        except Exception:
            pass
    return []


@router.get("/{company}")
def get_visualization(
    company: str,
    tipo1: str, año1_ini: int, mes1_ini: int, año1_fin: int, mes1_fin: int,
    tipo2: str, año2_ini: int, mes2_ini: int, año2_fin: int, mes2_fin: int,
):
    if not settings.parquet_path.exists():
        raise HTTPException(404, "Parquet no encontrado")

    df_all = pq.get_cached_df(settings.parquet_path)
    df_co  = df_all[df_all["compania"] == company]

    cols1 = _build_columns(año1_ini, mes1_ini, año1_fin, mes1_fin)
    cols2 = _build_columns(año2_ini, mes2_ini, año2_fin, mes2_fin)

    # Pre-build snapshots ONCE per tipo (not once per KPI)
    def _build_snaps(cols, año_fin, mes_fin, tipo):
        snaps = {}
        for yr in sorted({c["year"] for c in cols}):
            snap_mes = mes_fin if yr == año_fin else 12
            snaps[yr] = _fetch_year_snapshot(df_co, tipo, yr, snap_mes)
        return snaps

    snaps1 = _build_snaps(cols1, año1_fin, mes1_fin, tipo1)
    snaps2 = _build_snaps(cols2, año2_fin, mes2_fin, tipo2)

    company_items   = _get_company_items(company)
    mappings        = load_mappings().get(company, {})
    viz_scales      = _load_visualizar_scales(company)

    def resolve_sources(mapped_value) -> list[str]:
        if not mapped_value or mapped_value == "__NA__":
            return []
        if isinstance(mapped_value, dict):
            if mapped_value.get("_s"):
                src = mapped_value.get("src", "")
                return [src] if src else []
            return []  # _f formula — handled separately
        return [mapped_value] if isinstance(mapped_value, str) else [s for s in mapped_value if s]

    sections = []
    cur_sec = ""
    cur_subhdr = ""
    cur_kpis: list = []

    for item in company_items:
        if item["type"] == "header":
            if cur_sec and cur_kpis:
                sections.append({"label": cur_sec, "kpis": cur_kpis})
            cur_sec = item["label"]
            cur_subhdr = ""
            cur_kpis = []
            continue

        if item["type"] == "subheader":
            cur_subhdr = item["label"]
            cur_kpis.append({"type": "subheader", "label": item["label"], "is_na": False})
            continue

        label = item["label"]
        lk    = f"{cur_sec}||{cur_subhdr}||{label}" if cur_subhdr else f"{cur_sec}||{label}"
        mapped = mappings.get(lk)
        src1 = resolve_sources(mapped)
        src2 = resolve_sources(mapped)
        is_formula = isinstance(mapped, dict) and bool(mapped.get("_f"))
        is_scalar  = isinstance(mapped, dict) and bool(mapped.get("_s"))

        vals1 = _get_series(snaps1, cols1, src1)
        vals2 = _get_series(snaps2, cols2, src2)

        # Apply display scale from formula_params (e.g. fraction→percent for CEN ley)
        _vsc = viz_scales.get(lk)
        if _vsc is not None:
            vals1 = [v * _vsc if v is not None else None for v in vals1]
            vals2 = [v * _vsc if v is not None else None for v in vals2]

        if is_scalar and mapped.get("scalar") is not None:
            sc_op  = mapped.get("op", "*")
            sc_val = float(mapped["scalar"])
            def _sc(v, _op=sc_op, _val=sc_val): return _apply_op(v, _op, _val)
            vals1 = [_sc(v) for v in vals1]
            vals2 = [_sc(v) for v in vals2]

        mes1, ytd1 = _derive_mes_ytd(vals1, cols1, año1_fin)
        mes2, ytd2 = _derive_mes_ytd(vals2, cols2, año2_fin)

        cur_kpis.append({
            "label": label,
            "unit":  item.get("unit", ""),
            "lk":    lk,
            "mapped": bool(mapped and mapped != "__NA__"),
            "is_na": mapped == "__NA__",
            "is_formula": is_formula,
            "vals1": vals1, "mes1": mes1, "ytd1": ytd1,
            "vals2": vals2, "mes2": mes2, "ytd2": ytd2,
        })

    if cur_sec and cur_kpis:
        sections.append({"label": cur_sec, "kpis": cur_kpis})

    # Second pass: compute formula KPIs
    by_lk1: dict[str, list] = {}
    by_lk2: dict[str, list] = {}
    mes_by_lk1: dict[str, float | None] = {}
    mes_by_lk2: dict[str, float | None] = {}
    ytd_by_lk1: dict[str, float | None] = {}
    ytd_by_lk2: dict[str, float | None] = {}
    for sec in sections:
        for kpi in sec["kpis"]:
            if kpi.get("type") == "subheader":
                continue
            by_lk1[kpi["lk"]] = kpi["vals1"]
            by_lk2[kpi["lk"]] = kpi["vals2"]
            mes_by_lk1[kpi["lk"]] = kpi["mes1"]
            mes_by_lk2[kpi["lk"]] = kpi["mes2"]
            ytd_by_lk1[kpi["lk"]] = kpi["ytd1"]
            ytd_by_lk2[kpi["lk"]] = kpi["ytd2"]

    for sec in sections:
        for kpi in sec["kpis"]:
            if kpi.get("type") == "subheader":
                continue
            mapped = mappings.get(kpi["lk"])
            if not isinstance(mapped, dict):
                continue

            if mapped.get("_f"):
                a, op, b = mapped.get("a",""), mapped.get("op","/"), mapped.get("b","")
                a1 = by_lk1.get(a, _empty_series(len(cols1)))
                b1 = by_lk1.get(b, _empty_series(len(cols1)))
                a2 = by_lk2.get(a, _empty_series(len(cols2)))
                b2 = by_lk2.get(b, _empty_series(len(cols2)))
                kpi["vals1"] = [_apply_scale(_apply_op(x, op, y), mapped) for x, y in zip(a1, b1)]
                kpi["vals2"] = [_apply_scale(_apply_op(x, op, y), mapped) for x, y in zip(a2, b2)]
                kpi["mes1"]  = _apply_scale(_apply_op(mes_by_lk1.get(a), op, mes_by_lk1.get(b)), mapped)
                kpi["mes2"]  = _apply_scale(_apply_op(mes_by_lk2.get(a), op, mes_by_lk2.get(b)), mapped)
                kpi["ytd1"]  = _apply_scale(_apply_op(ytd_by_lk1.get(a), op, ytd_by_lk1.get(b)), mapped)
                kpi["ytd2"]  = _apply_scale(_apply_op(ytd_by_lk2.get(a), op, ytd_by_lk2.get(b)), mapped)
                kpi["mapped"] = bool(a and b)

            elif mapped.get("_pf"):
                num, den, op = mapped.get("num",""), mapped.get("den",""), mapped.get("op","/")
                def _pf_col(snaps, cols, _num=num, _den=den, _op=op):
                    out = []
                    for col in cols:
                        snap = snaps.get(col["year"], {})
                        f    = MONTH_BY_NUM[col["month"]]
                        nv   = _resolve_snapshot(snap, _num).get(f) if _num else None
                        dv   = _resolve_snapshot(snap, _den).get(f) if _den else None
                        out.append(_apply_scale(_apply_op(nv, _op, dv), mapped))
                    return out
                def _pf_agg(snaps, año_fin, cols, _num=num, _den=den, _op=op):
                    vals = _pf_col(snaps, cols)
                    return _derive_mes_ytd(vals, cols, año_fin)
                kpi["vals1"] = _pf_col(snaps1, cols1)
                kpi["vals2"] = _pf_col(snaps2, cols2)
                kpi["mes1"], kpi["ytd1"] = _derive_mes_ytd(kpi["vals1"], cols1, año1_fin)
                kpi["mes2"], kpi["ytd2"] = _derive_mes_ytd(kpi["vals2"], cols2, año2_fin)
                kpi["mapped"] = bool(num and den)

    return {
        "company": company,
        "tipo1": tipo1, "cols1": cols1,
        "tipo2": tipo2, "cols2": cols2,
        "sections": sections,
    }


# ── Excel export ─────────────────────────────────────────────────────────────

def _sf(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _build_excel(data: dict) -> io.BytesIO:
    company = data["company"]
    cols1, cols2 = data["cols1"], data["cols2"]
    tipo1, tipo2 = data["tipo1"], data["tipo2"]
    nc1, nc2 = len(cols1), len(cols2)

    # Column indices (1-based)
    C_KPI  = 1
    C_UNIT = 2
    C_TAG  = 3
    C_T1   = 4                        # first tipo1 month col
    C_MES1 = C_T1 + nc1
    C_YTD1 = C_MES1 + 1
    C_T2   = C_YTD1 + 1               # first tipo2 month col
    C_MES2 = C_T2 + nc2
    C_YTD2 = C_MES2 + 1
    LAST   = C_YTD2

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = company

    def cell(r, c, value=None, bold=False, color="000000", bg=None, center=False, right=False, size=9, italic=False):
        ce = ws.cell(row=r, column=c, value=value)
        ce.font = Font(bold=bold, color=color, size=size, italic=italic)
        if bg:
            ce.fill = _sf(bg)
        if center:
            ce.alignment = Alignment(horizontal="center", vertical="center")
        elif right:
            ce.alignment = Alignment(horizontal="right", vertical="center")
        else:
            ce.alignment = Alignment(vertical="center")
        return ce

    def num(r, c, value, bold=False):
        ce = ws.cell(row=r, column=c, value=round(value, 2) if value is not None else None)
        ce.font = Font(bold=bold, size=9)
        ce.alignment = Alignment(horizontal="right", vertical="center")
        ce.number_format = "#,##0.00"
        return ce

    # ── Row 1: group headers ────────────────────────────────────────────────
    ws.merge_cells(start_row=1, start_column=C_KPI, end_row=1, end_column=C_TAG)
    cell(1, C_KPI, company, bold=True, color="FFFFFF", bg="1A3A4A", center=True, size=10)

    ws.merge_cells(start_row=1, start_column=C_T1, end_row=1, end_column=C_YTD1)
    cell(1, C_T1, tipo1, bold=True, color="FFFFFF", bg="1A7080", center=True, size=10)

    ws.merge_cells(start_row=1, start_column=C_T2, end_row=1, end_column=C_YTD2)
    cell(1, C_T2, tipo2, bold=True, color="FFFFFF", bg="5a3c00", center=True, size=10)

    # ── Row 2: column labels ────────────────────────────────────────────────
    for c, lbl in [(C_KPI, "KPI"), (C_UNIT, "Ud."), (C_TAG, "Tipo")]:
        cell(2, c, lbl, bold=True, color="FFFFFF", bg="1A3A4A", center=True)

    for i, col in enumerate(cols1):
        cell(2, C_T1 + i, col["label"], bold=True, color="FFFFFF", bg="1A7080", center=True)
    cell(2, C_MES1, "Mes", bold=True, color="FFFFFF", bg="0e4e5e", center=True)
    cell(2, C_YTD1, "YTD", bold=True, color="FFFFFF", bg="0e4e5e", center=True)

    for i, col in enumerate(cols2):
        cell(2, C_T2 + i, col["label"], bold=True, color="FFFFFF", bg="5a3c00", center=True)
    cell(2, C_MES2, "Mes", bold=True, color="FFFFFF", bg="5a3c00", center=True)
    cell(2, C_YTD2, "YTD", bold=True, color="FFFFFF", bg="5a3c00", center=True)

    # ── Data rows ───────────────────────────────────────────────────────────
    row = 3
    for sec in data["sections"]:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=LAST)
        cell(row, 1, sec["label"], bold=True, color="FFFFFF", bg="1A3A4A", size=9)
        row += 1

        for kpi in sec["kpis"]:
            if kpi.get("type") == "subheader":
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=LAST)
                cell(row, 1, f"  {kpi['label']}", bold=True, color="FFFFFF", bg="2d6070", size=8)
                row += 1
                continue

            mapped  = kpi.get("mapped", False)
            txt_col = "000000" if mapped else "AAAAAA"

            # tipo1 row
            ws.merge_cells(start_row=row, start_column=C_KPI, end_row=row + 1, end_column=C_KPI)
            ws.merge_cells(start_row=row, start_column=C_UNIT, end_row=row + 1, end_column=C_UNIT)
            kpi_cell = ws.cell(row=row, column=C_KPI, value=kpi["label"])
            kpi_cell.font = Font(color=txt_col, size=9, italic=not mapped)
            kpi_cell.alignment = Alignment(vertical="center")
            unit_cell = ws.cell(row=row, column=C_UNIT, value=kpi.get("unit", ""))
            unit_cell.font = Font(color="888888", size=8)
            unit_cell.alignment = Alignment(horizontal="center", vertical="center")

            cell(row, C_TAG, tipo1, color="1A7080", bg="e0f2f1", center=True, size=8)
            for i, v in enumerate(kpi["vals1"]):
                num(row, C_T1 + i, v)
            num(row, C_MES1, kpi["mes1"], bold=True)
            num(row, C_YTD1, kpi["ytd1"], bold=True)
            row += 1

            # tipo2 row
            cell(row, C_TAG, tipo2, color="92400e", bg="fef3c7", center=True, size=8)
            for i, v in enumerate(kpi["vals2"]):
                num(row, C_T2 + i, v)
            num(row, C_MES2, kpi["mes2"], bold=True)
            num(row, C_YTD2, kpi["ytd2"], bold=True)
            # bottom border between KPIs
            for c in range(1, LAST + 1):
                ws.cell(row=row, column=c).border = openpyxl.styles.Border(
                    bottom=openpyxl.styles.Side(style="thin", color="DDDDDD")
                )
            row += 1

    # ── Column widths & freeze ──────────────────────────────────────────────
    cl = openpyxl.utils.get_column_letter
    ws.column_dimensions[cl(C_KPI)].width  = 30
    ws.column_dimensions[cl(C_UNIT)].width = 7
    ws.column_dimensions[cl(C_TAG)].width  = 9
    for c in range(C_T1, LAST + 1):
        ws.column_dimensions[cl(c)].width = 11
    ws.freeze_panes = "D3"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@router.get("/{company}/excel")
def export_excel(
    company: str,
    tipo1: str, año1_ini: int, mes1_ini: int, año1_fin: int, mes1_fin: int,
    tipo2: str, año2_ini: int, mes2_ini: int, año2_fin: int, mes2_fin: int,
):
    data = get_visualization(
        company, tipo1, año1_ini, mes1_ini, año1_fin, mes1_fin,
        tipo2, año2_ini, mes2_ini, año2_fin, mes2_fin,
    )
    buf = _build_excel(data)
    filename = f"Input_{company}_{tipo1}_{año1_fin}{mes1_fin:02d}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
