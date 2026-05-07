"""
Reporte Excel completo: todas las compañías + Grupo, todas las secciones.
  Hoja 1 «Input»          — KPIs visualizados (tipo1 / tipo2, columnas mensuales)
  Hoja 2 «Costos Aj»      — Costos Ajustados (Mes + YTD); Grupo con fórmulas SUM
  Hoja 3 «Cálculos»       — Descomposición de Varianzas; Mes y YTD como fórmulas Excel
  Hoja 4 «Análisis Costo» — Waterfall por efecto; Grupo con fórmulas SUM
  Hoja 5 «Explicaciones»  — Paneles Real/Ppto/Efecto (Mes + YTD)
"""
import io
from fastapi import APIRouter
from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

from .visualizar       import get_visualization
from .costos_ajustados import get_costos_ajustados
from .calculos         import get_calculos
from .explicaciones    import get_explicaciones

router = APIRouter(prefix="/api/export", tags=["export"])

COMPANIES = ["MLP", "CEN", "ANT", "CMZ"]

# ── Colores ──────────────────────────────────────────────────────────────────
NAVY  = "1A3A4A"; TEAL  = "1A7080"; TEAL2 = "0e4e5e"
AMBER = "5a3c00"; GRN1  = "065f46"; GRN2  = "047857"
GRN3  = "059669"; BLUE  = "2d6070"; WHT   = "FFFFFF"

# ── Helpers de celda ─────────────────────────────────────────────────────────

def _sf(c): return PatternFill("solid", fgColor=c)
def _cl(c): return openpyxl.utils.get_column_letter(c)

def _wc(ws, r, c, val=None, bold=False, color="000000", bg=None, h="left", size=9, italic=False):
    ce = ws.cell(row=r, column=c, value=val)
    ce.font = Font(bold=bold, color=color, size=size, italic=italic)
    if bg: ce.fill = _sf(bg)
    ce.alignment = Alignment(horizontal=h, vertical="center")
    return ce

def _num(ws, r, c, val, bold=False):
    """Escribe un valor numérico literal."""
    if val is not None:
        val = round(float(val), 2)
    ce = ws.cell(row=r, column=c, value=val)
    ce.font = Font(bold=bold, size=9)
    ce.alignment = Alignment(horizontal="right", vertical="center")
    ce.number_format = "#,##0.00"

def _fml(ws, r, c, formula: str, bold=False):
    """Escribe una fórmula Excel (cadena que empieza con =)."""
    ce = ws.cell(row=r, column=c, value=formula)
    ce.font = Font(bold=bold, size=9)
    ce.alignment = Alignment(horizontal="right", vertical="center")
    ce.number_format = "#,##0.00"

def _hdr(ws, r, c1, c2, val, bg, color=WHT, bold=True, size=9, h="left"):
    if c1 < c2:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    ce = ws.cell(row=r, column=c1, value=val)
    ce.font = Font(bold=bold, color=color, size=size)
    ce.fill = _sf(bg)
    ce.alignment = Alignment(horizontal=h, vertical="center")

def _co_hdr(ws, r, c2, company):
    _hdr(ws, r, 1, c2, company, BLUE, WHT, bold=True, size=10)

def _sec_hdr(ws, r, c2, label):
    _hdr(ws, r, 1, c2, label, NAVY, WHT, bold=True, size=9)

def _try(fn, *a, **kw):
    try:
        return fn(*a, **kw)
    except Exception:
        return None

def _ref(c, r):
    """Referencia absoluta tipo $A$1."""
    return f"${_cl(c)}${r}"

def _sum_refs(refs: list[str]) -> str:
    """Fórmula =refs[0]+refs[1]+... (None ignorados)."""
    valid = [r for r in refs if r]
    if not valid:
        return ""
    return "=" + "+".join(valid)


# ── Hoja 1: Input ────────────────────────────────────────────────────────────

def _input_sheet(wb, tipo1, a1i, m1i, a1f, m1f, tipo2, a2i, m2i, a2f, m2f):
    ws = wb.create_sheet("Input")

    all_data = [_try(get_visualization, co, tipo1, a1i, m1i, a1f, m1f, tipo2, a2i, m2i, a2f, m2f)
                for co in COMPANIES]
    first = next((d for d in all_data if d), None)
    if not first:
        ws.cell(1, 1, "Sin datos"); return

    cols1, cols2 = first["cols1"], first["cols2"]
    nc = max(len(cols1), len(cols2))

    # Columnas de meses compartidas entre Real y Ppto (igual que el frontend)
    C_KPI = 1; C_UD = 2; C_TAG = 3
    C_MON = 4           # primer mes (compartido)
    C_MES = C_MON + nc  # columna Mes
    C_YTD = C_MES + 1   # columna YTD
    LAST  = C_YTD

    # Header row 1 (Real) — KPI y Ud. se fusionan con row 2
    ws.merge_cells(start_row=1, start_column=C_KPI, end_row=2, end_column=C_KPI)
    ce = ws.cell(row=1, column=C_KPI, value="KPI")
    ce.font = Font(bold=True, color=WHT, size=9)
    ce.fill = _sf(NAVY)
    ce.alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells(start_row=1, start_column=C_UD, end_row=2, end_column=C_UD)
    ce = ws.cell(row=1, column=C_UD, value="Ud.")
    ce.font = Font(bold=True, color=WHT, size=9)
    ce.fill = _sf(NAVY)
    ce.alignment = Alignment(horizontal="center", vertical="center")

    _wc(ws, 1, C_TAG, tipo1, bold=True, color=WHT, bg=TEAL, h="center")
    for i, col in enumerate(cols1):
        _wc(ws, 1, C_MON+i, col["label"], bold=True, color=WHT, bg=TEAL, h="center")
    _wc(ws, 1, C_MES, "Mes", bold=True, color=WHT, bg=TEAL2, h="center")
    _wc(ws, 1, C_YTD, "YTD", bold=True, color=WHT, bg=TEAL2, h="center")

    # Header row 2 (Ppto)
    _wc(ws, 2, C_TAG, tipo2, bold=True, color=WHT, bg=AMBER, h="center")
    for i, col in enumerate(cols2):
        _wc(ws, 2, C_MON+i, col["label"], bold=True, color=WHT, bg=AMBER, h="center")
    _wc(ws, 2, C_MES, "Mes", bold=True, color=WHT, bg=AMBER, h="center")
    _wc(ws, 2, C_YTD, "YTD", bold=True, color=WHT, bg=AMBER, h="center")

    row = 3
    for co, data in zip(COMPANIES, all_data):
        if not data: continue
        _co_hdr(ws, row, LAST, co); row += 1

        for sec in data["sections"]:
            _sec_hdr(ws, row, LAST, sec["label"]); row += 1
            for kpi in sec["kpis"]:
                if kpi.get("type") == "subheader":
                    _hdr(ws, row, 1, LAST, f"  {kpi['label']}", BLUE, WHT, size=8); row += 1
                    continue
                mapped = kpi.get("mapped", False)
                tcol   = "000000" if mapped else "AAAAAA"

                # Fusionar KPI y Ud. entre fila Real y fila Ppto
                ws.merge_cells(start_row=row, start_column=C_KPI, end_row=row+1, end_column=C_KPI)
                ws.merge_cells(start_row=row, start_column=C_UD,  end_row=row+1, end_column=C_UD)

                ce = ws.cell(row=row, column=C_KPI, value=kpi["label"])
                ce.font = Font(color=tcol, size=9, italic=not mapped)
                ce.alignment = Alignment(vertical="center")

                ce2 = ws.cell(row=row, column=C_UD, value=kpi.get("unit", ""))
                ce2.font = Font(color="888888", size=8)
                ce2.alignment = Alignment(horizontal="center", vertical="center")

                # Fila Real (tipo1)
                _wc(ws, row, C_TAG, tipo1, bold=True, color=WHT, bg=TEAL, h="center", size=8)
                for i, v in enumerate(kpi["vals1"]):
                    _num(ws, row, C_MON+i, v)
                _num(ws, row, C_MES, kpi["mes1"], bold=True)
                _num(ws, row, C_YTD, kpi["ytd1"], bold=True)

                # Fila Ppto (tipo2)
                _wc(ws, row+1, C_TAG, tipo2, bold=True, color=WHT, bg=AMBER, h="center", size=8)
                for i, v in enumerate(kpi["vals2"]):
                    _num(ws, row+1, C_MON+i, v)
                _num(ws, row+1, C_MES, kpi["mes2"], bold=True)
                _num(ws, row+1, C_YTD, kpi["ytd2"], bold=True)

                row += 2
        row += 1

    ws.column_dimensions[_cl(C_KPI)].width = 28
    ws.column_dimensions[_cl(C_UD)].width  = 6
    ws.column_dimensions[_cl(C_TAG)].width = 8
    for c in range(C_MON, LAST+1):
        ws.column_dimensions[_cl(c)].width = 11
    ws.freeze_panes = "D3"


# ── Hoja 2: Costos Ajustados — Grupo con fórmulas SUM ────────────────────────

def _costos_aj_sheet(wb, año, mes):
    ws = wb.create_sheet("Costos Aj")

    C_LBL = 1
    C_MR=2; C_MP=3; C_MPA=4; C_MD=5
    C_YR=6; C_YP=7; C_YPA=8; C_YD=9
    LAST = 9
    DATA_COLS = [C_MR, C_MP, C_MPA, C_MD, C_YR, C_YP, C_YPA, C_YD]

    _hdr(ws, 1, 1, 1,    "KPI",  NAVY, WHT)
    _hdr(ws, 1, 2, 5,    "Mes",  TEAL, WHT, h="center")
    _hdr(ws, 1, 6, 9,    "YTD",  AMBER, WHT, h="center")
    for c, lbl, bg in [(1,"KPI",NAVY),(2,"Real",TEAL),(3,"Ppto",TEAL),(4,"PptoAj",TEAL),(5,"Dif",TEAL),
                        (6,"Real",AMBER),(7,"Ppto",AMBER),(8,"PptoAj",AMBER),(9,"Dif",AMBER)]:
        _wc(ws, 2, c, lbl, bold=True, color=WHT, bg=bg, h="center")

    row = 3
    KEYS = [("mes","real"),("mes","ppto"),("mes","ppto_aj_tot"),("mes","dif"),
            ("ytd","real"),("ytd","ppto"),("ytd","ppto_aj_tot"),("ytd","dif")]
    TOT_STYLE = [
        ("costo_onsite","Costo Onsite",NAVY),("pre_credito","Pre-crédito",BLUE),
        ("c1","C1",GRN1),("c2","C2",GRN2),("c3","C3",GRN3),
    ]

    def _fill_row(r, item_data, bold=False):
        for col, (period, key) in zip(DATA_COLS, KEYS):
            _num(ws, r, col, (item_data.get(period) or {}).get(key), bold=bold)

    # dict: tot_key -> {col -> [row_number per company]}
    total_refs: dict[str, dict[int, list[int]]] = {}

    for co in COMPANIES:
        d = _try(get_costos_ajustados, co, año, mes)
        if not d: continue
        _co_hdr(ws, row, LAST, co); row += 1

        for grp in d.get("grupos", []):
            gt = grp.get("total", {})
            _wc(ws, row, C_LBL, grp["label"], bold=True, color=WHT, bg=NAVY, size=9)
            _fill_row(row, gt, bold=True); row += 1
            sas = grp.get("subareas", [])
            if len(sas) > 1:
                for sa in sas:
                    if not sa.get("mapped", False): continue
                    _wc(ws, row, C_LBL, f"  {sa['label']}", italic=True, size=8)
                    _fill_row(row, sa); row += 1

        for sr in d.get("summary", []):
            _wc(ws, row, C_LBL, sr["label"], bold=True, bg="fef3c7", size=9)
            _fill_row(row, sr, bold=True); row += 1

        # Totales por compañía — registrar fila para fórmulas de Grupo
        for key, label, bg in TOT_STYLE:
            tot = d.get("totales", {}).get(key, {})
            _wc(ws, row, C_LBL, label, bold=True, color=WHT, bg=bg, size=9)
            _fill_row(row, tot, bold=True)
            # Registrar referencia de cada columna de datos
            total_refs.setdefault(key, {})
            for col in DATA_COLS:
                total_refs[key].setdefault(col, []).append(row)
            row += 1
        row += 1

    # Grupo Minero — fórmulas SUM referenciando filas de compañías
    if total_refs:
        _co_hdr(ws, row, LAST, "Grupo Minero"); row += 1
        for key, label, bg in TOT_STYLE:
            _wc(ws, row, C_LBL, label, bold=True, color=WHT, bg=bg, size=9)
            refs = total_refs.get(key, {})
            for col in DATA_COLS:
                row_list = refs.get(col, [])
                if row_list:
                    formula = "=" + "+".join(f"{_cl(col)}{r}" for r in row_list)
                    _fml(ws, row, col, formula, bold=True)
            row += 1

    ws.column_dimensions[_cl(C_LBL)].width = 32
    for c in range(2, LAST+1):
        ws.column_dimensions[_cl(c)].width = 12
    ws.freeze_panes = "B3"


# ── Hoja 3: Cálculos — Mes y YTD como fórmulas Excel ────────────────────────

_CALC_SECS = [
    ("Delta Gasto Precio", "delta_gasto_precio", "kUS$", [
        ("Bolas","bolas"),("Ácido","acido"),("Energía","energia"),
        ("Combustible","combustible"),("Explosivos","explosivos"),("TOTAL","total"),
    ]),
    ("Delta Rendimiento", "delta_rendimiento", "kUS$", [
        ("Bolas","bolas"),("Ácido","acido"),("Energía conc","energia_conc"),
        ("Energía hidro","energia_hidro"),("Combustible","combustible"),
        ("Explosivos","explosivos"),("TOTAL","total"),
    ]),
    ("Delta Actividad", "delta_actividad", "CU", [
        ("Trat. Conc","trat_concentradora"),("Trat. Hidro","trat_hidro"),
        ("Rec. Conc","rec_concentradora"),("Rec. Hidro","rec_hidro"),("TOTAL","total"),
    ]),
    ("Efecto Ley", "efecto_ley", "CU", [
        ("Concentradora","concentradora"),("Hidro","hidro"),("TOTAL","total"),
    ]),
    ("TC/RC", "efectos_tcrc", "kUS$", [
        ("Cantidades","cantidades"),("Tarifas","tarifas"),("TOTAL","total"),
    ]),
    ("Comercialización", "efectos_comer", "kUS$", [
        ("Cantidades","cantidades"),("Tarifas","tarifas"),("TOTAL","total"),
    ]),
]


def _calculos_sheet(wb, a_ini, m_ini, a_fin, m_fin):
    ws = wb.create_sheet("Cálculos")

    first_cols = None
    for co in COMPANIES:
        d = _try(get_calculos, co, a_ini, m_ini, a_fin, m_fin)
        if d: first_cols = d["cols"]; break
    if not first_cols:
        ws.cell(1, 1, "Sin datos"); return

    nc = len(first_cols)
    C_COMP = 1; C_UD = 2; C_M_START = 3
    C_MES = C_M_START + nc      # columna Mes = SUM o referencia
    C_YTD = C_MES + 1           # columna YTD = SUM meses
    LAST  = C_YTD
    C_LAST_MON = C_MES - 1      # última columna de mes (= C_M_START + nc - 1)

    # Fila 1: cabeceras de grupo
    _hdr(ws, 1, C_COMP, C_UD,           "Componente",  NAVY, WHT)
    _hdr(ws, 1, C_M_START, C_LAST_MON,  "Meses",       TEAL, WHT, h="center")
    _wc(ws, 1, C_MES, "Mes (último)",   bold=True, color=WHT, bg=TEAL2, h="center")
    _wc(ws, 1, C_YTD, "YTD (acum.)",   bold=True, color=WHT, bg=TEAL2, h="center")

    # Fila 2: etiquetas individuales
    _wc(ws, 2, C_COMP, "Componente", bold=True, color=WHT, bg=NAVY, h="center")
    _wc(ws, 2, C_UD,   "Ud.",        bold=True, color=WHT, bg=NAVY, h="center")
    for i, col in enumerate(first_cols):
        _wc(ws, 2, C_M_START+i, col["label"], bold=True, color=WHT, bg=TEAL, h="center")
    _wc(ws, 2, C_MES, "Mes", bold=True, color=WHT, bg=TEAL2, h="center")
    _wc(ws, 2, C_YTD, "YTD", bold=True, color=WHT, bg=TEAL2, h="center")

    row = 3
    # Para Grupo: acumular referencias de filas TOTAL por sección
    grupo_tot_refs: dict[str, list[int]] = {}   # sec_key -> [row de TOTAL por compañía]

    for co in COMPANIES:
        d = _try(get_calculos, co, a_ini, m_ini, a_fin, m_fin)
        if not d: continue

        _co_hdr(ws, row, LAST, co); row += 1

        for sec_lbl, sec_key, unit, rows_def in _CALC_SECS:
            _sec_hdr(ws, row, LAST, sec_lbl); row += 1
            sec = d.get(sec_key, {})
            for r_lbl, r_key in rows_def:
                series  = sec.get(r_key, {})
                is_tot  = r_key == "total"
                vals    = series.get("vals", [None] * nc)

                _wc(ws, row, C_COMP, f"  {r_lbl}", bold=is_tot, size=9)
                _wc(ws, row, C_UD,   unit, color="888888", size=8, h="center")

                # Columnas mensuales: valores literales
                for i, v in enumerate(vals):
                    _num(ws, row, C_M_START+i, v, bold=is_tot)

                # Mes = referencia a la última columna mensual de esta fila
                _fml(ws, row, C_MES,
                     f"={_cl(C_LAST_MON)}{row}",
                     bold=is_tot)

                # YTD = SUM de todas las columnas mensuales de esta fila
                _fml(ws, row, C_YTD,
                     f"=SUM({_cl(C_M_START)}{row}:{_cl(C_LAST_MON)}{row})",
                     bold=is_tot)

                if is_tot:
                    grupo_tot_refs.setdefault(sec_key, []).append(row)
                row += 1
        row += 1

    # Grupo Minero — fórmulas que suman las filas TOTAL de cada compañía
    if grupo_tot_refs:
        _co_hdr(ws, row, LAST, "Grupo Minero"); row += 1
        for sec_lbl, sec_key, unit, _ in _CALC_SECS:
            refs = grupo_tot_refs.get(sec_key, [])
            if not refs: continue
            _sec_hdr(ws, row, LAST, sec_lbl); row += 1
            _wc(ws, row, C_COMP, "  TOTAL", bold=True, size=9)
            _wc(ws, row, C_UD, unit, color="888888", size=8, h="center")
            # Columnas mensuales: suma de los TOTAL de cada compañía
            for i in range(nc):
                col = C_M_START + i
                _fml(ws, row, col,
                     "=" + "+".join(f"{_cl(col)}{r}" for r in refs),
                     bold=True)
            # Mes = referencia a la última columna mensual del grupo (ya es fórmula)
            _fml(ws, row, C_MES,
                 f"={_cl(C_LAST_MON)}{row}",
                 bold=True)
            # YTD = SUM del rango mensual del grupo
            _fml(ws, row, C_YTD,
                 f"=SUM({_cl(C_M_START)}{row}:{_cl(C_LAST_MON)}{row})",
                 bold=True)
            row += 1
        row += 1

    ws.column_dimensions[_cl(C_COMP)].width = 24
    ws.column_dimensions[_cl(C_UD)].width   = 7
    for c in range(C_M_START, LAST+1):
        ws.column_dimensions[_cl(c)].width = 11
    ws.freeze_panes = "C3"


# ── Hoja 4: Análisis Costo — Grupo con fórmulas SUM ─────────────────────────

_WF_SECS = [
    ("Ajuste Monetario (kUS$)", [
        ("aj_tc","TC Dólar"),("aj_ipc","IPC"),("aj_cpi","CPI"),
    ]),
    ("Precio Insumos (kUS$)", [
        ("pr_energia","Energía"),("pr_combustible","Combustible"),
        ("pr_acido","Ácido"),("pr_bolas","Bolas"),("pr_explosivos","Explosivos"),
    ]),
    ("Delta Actividad (kUS$)", [
        ("act_mina_kus","Mina"),("act_conc_kus","Concentradora"),
    ]),
    ("Rendimiento Insumos (kUS$)", [
        ("rend_energia_conc","Energía Conc"),("rend_combustible","Combustible"),
        ("rend_bolas","Bolas"),("rend_explosivos","Explosivos"),
    ]),
    ("Producción (ktCuf)", [
        ("da_trat_conc","ΔTrat Conc"),("da_rec_conc","ΔRec Conc"),
        ("da_trat_hidro","ΔTrat Hidro"),("da_rec_hidro","ΔRec Hidro"),
        ("ley_conc","Efecto Ley Conc"),("ley_hidro","Efecto Ley Hidro"),
    ]),
    ("Otras Variaciones (kUS$)", [
        ("dev_mina","Desarrollo Mina"),("tcrc_kus","TC/RC"),
        ("comer_kus","Comercialización"),
    ]),
    ("Estructura Real (kUS$)", [
        ("c1_b","C1 Ppto"),("c1_r","C1 Real"),
        ("c3_b","C3 Ppto"),("c3_r","C3 Real"),
    ]),
]


def _analisis_sheet(wb, año, mes):
    ws = wb.create_sheet("Análisis Costo")

    C_LBL = 1; C_M = 2; C_Y = 3; LAST = 3
    _hdr(ws, 1, 1, 1,   "Efecto",    NAVY, WHT)
    _wc(ws, 1, C_M,     "Mes kUS$",  bold=True, color=WHT, bg=TEAL,  h="center")
    _wc(ws, 1, C_Y,     "YTD kUS$",  bold=True, color=WHT, bg=AMBER, h="center")

    row = 2
    # wf_key -> {"mes": [row...], "ytd": [row...]}
    grupo_refs: dict[str, dict[str, list[int]]] = {}

    for co in COMPANIES:
        d = _try(get_explicaciones, co, año, mes)
        if not d: continue

        _co_hdr(ws, row, LAST, co); row += 1
        wf_m = (d.get("data") or {}).get("mes", {}).get("waterfall", {})
        wf_y = (d.get("data") or {}).get("ytd", {}).get("waterfall", {})

        for sec_lbl, items in _WF_SECS:
            _sec_hdr(ws, row, LAST, sec_lbl); row += 1
            for key, lbl in items:
                _wc(ws, row, C_LBL, f"  {lbl}", size=9)
                _num(ws, row, C_M, wf_m.get(key))
                _num(ws, row, C_Y, wf_y.get(key))
                # Registrar fila para el Grupo
                grupo_refs.setdefault(key, {"mes": [], "ytd": []})
                grupo_refs[key]["mes"].append(row)
                grupo_refs[key]["ytd"].append(row)
                row += 1
        row += 1

    # Grupo Minero — fórmulas SUM apuntando a filas de compañías
    if grupo_refs:
        _co_hdr(ws, row, LAST, "Grupo Minero"); row += 1
        for sec_lbl, items in _WF_SECS:
            _sec_hdr(ws, row, LAST, sec_lbl); row += 1
            for key, lbl in items:
                _wc(ws, row, C_LBL, f"  {lbl}", size=9)
                refs = grupo_refs.get(key, {})
                mes_rows = refs.get("mes", [])
                ytd_rows = refs.get("ytd", [])
                if mes_rows:
                    _fml(ws, row, C_M,
                         "=" + "+".join(f"{_cl(C_M)}{r}" for r in mes_rows),
                         bold=True)
                if ytd_rows:
                    _fml(ws, row, C_Y,
                         "=" + "+".join(f"{_cl(C_Y)}{r}" for r in ytd_rows),
                         bold=True)
                row += 1

    ws.column_dimensions[_cl(C_LBL)].width = 28
    ws.column_dimensions[_cl(C_M)].width   = 14
    ws.column_dimensions[_cl(C_Y)].width   = 14
    ws.freeze_panes = "B2"


# ── Hoja 5: Explicaciones ────────────────────────────────────────────────────

def _explicaciones_sheet(wb, año, mes):
    ws = wb.create_sheet("Explicaciones")

    C_LBL = 1; C_UD = 2; C_R = 3; C_P = 4; C_E = 5; LAST = 5
    _wc(ws, 1, C_LBL, "Indicador",  bold=True, color=WHT, bg=NAVY)
    _wc(ws, 1, C_UD,  "Ud.",        bold=True, color=WHT, bg=NAVY, h="center")
    _wc(ws, 1, C_R,   "Real",       bold=True, color=WHT, bg=TEAL, h="center")
    _wc(ws, 1, C_P,   "Ppto",       bold=True, color=WHT, bg=TEAL, h="center")
    _wc(ws, 1, C_E,   "Efecto",     bold=True, color=WHT, bg=TEAL2, h="center")

    row = 2

    def _item_rows(ws, r, items, ef_key="efecto_kus"):
        for item in items:
            _wc(ws, r, C_LBL, f"  {item['label']}", size=9)
            _wc(ws, r, C_UD,  item.get("unit",""), color="888888", size=8, h="center")
            _num(ws, r, C_R,  item.get("real"))
            _num(ws, r, C_P,  item.get("ppto"))
            _num(ws, r, C_E,  item.get(ef_key), bold=True)
            r += 1
        return r

    for co in COMPANIES:
        d = _try(get_explicaciones, co, año, mes)
        if not d: continue
        for panel_name in ("mes", "ytd"):
            panel = (d.get("data") or {}).get(panel_name, {})
            _co_hdr(ws, row, LAST, f"{co} — {panel_name.upper()}"); row += 1

            _sec_hdr(ws, row, LAST, "Ajuste Monetario"); row += 1
            row = _item_rows(ws, row, panel.get("ajuste_monetario", []))

            _sec_hdr(ws, row, LAST, "Precio Insumos"); row += 1
            row = _item_rows(ws, row, panel.get("precio_insumos", []))

            _sec_hdr(ws, row, LAST, "Rendimiento Insumos"); row += 1
            row = _item_rows(ws, row, panel.get("rendimiento", []))

            _sec_hdr(ws, row, LAST, "Producción"); row += 1
            row = _item_rows(ws, row, panel.get("produccion", []), ef_key="efecto_ktcuf")

            _sec_hdr(ws, row, LAST, "Mov. Mina / Desarrollo Mina"); row += 1
            mv = panel.get("mov_mina", {})
            _wc(ws, row, C_LBL, "  Mov. Mina", size=9)
            _num(ws, row, C_R, mv.get("real"))
            _num(ws, row, C_P, mv.get("ppto"))
            _num(ws, row, C_E, mv.get("efecto_kus"), bold=True)
            row += 1
            dm = panel.get("desarrollo_mina", {})
            _wc(ws, row, C_LBL, "  Desarrollo Mina", size=9)
            _num(ws, row, C_R, dm.get("real"))
            _num(ws, row, C_P, dm.get("ppto"))
            _num(ws, row, C_E, dm.get("efecto_kus"), bold=True)
            row += 1
            row += 1
        row += 1

    ws.column_dimensions[_cl(C_LBL)].width = 28
    ws.column_dimensions[_cl(C_UD)].width  = 8
    for c in range(C_R, LAST+1):
        ws.column_dimensions[_cl(c)].width = 13
    ws.freeze_panes = "C2"


# ── Endpoint ─────────────────────────────────────────────────────────────────

@router.get("/excel")
def export_full_excel(
    tipo1: str, año1_ini: int, mes1_ini: int, año1_fin: int, mes1_fin: int,
    tipo2: str, año2_ini: int, mes2_ini: int, año2_fin: int, mes2_fin: int,
):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _input_sheet(wb, tipo1, año1_ini, mes1_ini, año1_fin, mes1_fin,
                     tipo2, año2_ini, mes2_ini, año2_fin, mes2_fin)
    _costos_aj_sheet(wb, año1_fin, mes1_fin)
    _calculos_sheet(wb, año1_ini, mes1_ini, año1_fin, mes1_fin)
    _analisis_sheet(wb, año1_fin, mes1_fin)
    _explicaciones_sheet(wb, año1_fin, mes1_fin)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fn = f"Reporte_Robot_{año1_fin}{mes1_fin:02d}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fn}"'},
    )
