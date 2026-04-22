import shutil
from pathlib import Path
from datetime import datetime

import openpyxl
import pandas as pd

from .mapper_service import (
    COMPANY_SHEET,
    MONTH_COL,
    MES_COL,
    YTD_COL,
    discover_sheet_structure,
    find_excel_row,
)
from .mapping_store import get_company_mappings
from .parquet_service import MONTH_FIELDS


MONTH_NUMBER = {field: i + 1 for i, field in enumerate(MONTH_FIELDS)}


def _write_value(ws, row: int, col: int, value):
    if value is not None:
        ws.cell(row=row, column=col, value=value)


def _update_periodo_sheet(wb: openpyxl.Workbook, anio: int, mes: int):
    """Updates the Periodo sheet with the selected year and month."""
    if "Periodo" not in wb.sheetnames:
        return
    ws = wb["Periodo"]
    mes_str = f"{mes:02d}"
    periodo_str = f"{anio}.{mes_str}"
    # Row 3 col C = "Mes actualizar"
    ws["C3"] = periodo_str


def _apply_op(a, op, b):
    if a is None or b is None:
        return None
    if op == "/" and b == 0:
        return None
    return {"+": a+b, "-": a-b, "*": a*b, "/": a/b}.get(op)


def _apply_scale(val, mapped: dict):
    if val is None:
        return None
    scale = mapped.get("scale")
    if scale is None:
        return val
    scale_op = mapped.get("scale_op", "*")
    return _apply_op(val, scale_op, float(scale))


def _write_formula_kpis(
    sheet_mapping: dict[str, int],
    company_mappings: dict,
    row_vals: dict[int, dict[int, float]],
    mes_limit: int,
) -> dict[int, dict[int, float]]:
    """Computes formula KPIs from already-accumulated row values."""
    from .mapper_service import _normalise
    extra: dict[int, dict[int, float]] = {}
    for lk, mapped in company_mappings.items():
        if not isinstance(mapped, dict) or not mapped.get("_f"):
            continue
        a_lk = mapped.get("a", "")
        b_lk = mapped.get("b", "")
        op   = mapped.get("op", "/")
        f_row = sheet_mapping.get(_normalise(lk))
        a_row = sheet_mapping.get(_normalise(a_lk))
        b_row = sheet_mapping.get(_normalise(b_lk))
        if not all([f_row, a_row, b_row]):
            continue
        for m in range(1, mes_limit + 1):
            col = MONTH_COL[m]
            res = _apply_scale(_apply_op(row_vals.get(a_row, {}).get(col), op, row_vals.get(b_row, {}).get(col)), mapped)
            if res is not None:
                extra.setdefault(f_row, {})[col] = res
        for col in [MES_COL, YTD_COL]:
            res = _apply_scale(_apply_op(row_vals.get(a_row, {}).get(col), op, row_vals.get(b_row, {}).get(col)), mapped)
            if res is not None:
                extra.setdefault(f_row, {})[col] = res
    return extra


def process_company(
    ws,
    df_company: pd.DataFrame,
    sheet_mapping: dict[str, int],
    mes_limit: int,
    company_mappings: dict | None = None,
) -> tuple[int, list[str]]:
    """
    Writes monthly values for all matched KPIs.
    Values for the same Excel row are SUMMED (supports multi-source mappings).
    Returns (kpis_written, unmatched_labels).
    """
    accum: dict[int, dict[int, float]] = {}
    unmatched: list[str] = []

    for _, parquet_row in df_company.iterrows():
        excel_row = find_excel_row(parquet_row, sheet_mapping, company_mappings)
        if excel_row is None:
            label = parquet_row.get("kpi") or parquet_row.get("subcategory", "")
            unmatched.append(str(label))
            continue

        if excel_row not in accum:
            accum[excel_row] = {}

        for m in range(1, mes_limit + 1):
            field = MONTH_FIELDS[m - 1]
            val = parquet_row.get(field)
            if pd.notna(val):
                col = MONTH_COL[m]
                accum[excel_row][col] = accum[excel_row].get(col, 0.0) + float(val)

        for col, field in [(MES_COL, "mes"), (YTD_COL, "ytd")]:
            val = parquet_row.get(field)
            if pd.notna(val):
                accum[excel_row][col] = accum[excel_row].get(col, 0.0) + float(val)

    # Write regular KPIs
    for excel_row, cols in accum.items():
        for col, value in cols.items():
            _write_value(ws, excel_row, col, value)

    # Compute and write formula KPIs
    if company_mappings:
        formula_vals = _write_formula_kpis(sheet_mapping, company_mappings, accum, mes_limit)
        for excel_row, cols in formula_vals.items():
            for col, value in cols.items():
                _write_value(ws, excel_row, col, value)

    return len(accum), unmatched


def run_automation(
    template_path: Path,
    output_dir: Path,
    df_filtered: pd.DataFrame,
    anio: int,
    mes: int,
    companies: list[str],
) -> tuple[Path, list[dict]]:
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = output_dir / f"Robot_{anio}_{mes:02d}_{timestamp}.xlsx"

    shutil.copy2(template_path, output_path)
    wb = openpyxl.load_workbook(output_path)

    _update_periodo_sheet(wb, anio, mes)

    company_results = []
    for company in companies:
        sheet_name = COMPANY_SHEET.get(company)
        if not sheet_name or sheet_name not in wb.sheetnames:
            company_results.append(
                {
                    "compania": company,
                    "kpis_escritos": 0,
                    "kpis_no_mapeados": 0,
                    "kpis_no_mapeados_detalle": [],
                    "ok": False,
                    "error": f"Hoja '{sheet_name}' no encontrada en el Excel",
                }
            )
            continue

        ws = wb[sheet_name]
        sheet_mapping = discover_sheet_structure(wb, company)
        df_company = df_filtered[df_filtered["compania"] == company]

        company_mappings = get_company_mappings(company)
        written, unmatched = process_company(ws, df_company, sheet_mapping, mes, company_mappings)
        company_results.append(
            {
                "compania": company,
                "kpis_escritos": written,
                "kpis_no_mapeados": len(unmatched),
                "kpis_no_mapeados_detalle": unmatched,
                "ok": True,
            }
        )

    wb.save(output_path)
    return output_path, company_results
